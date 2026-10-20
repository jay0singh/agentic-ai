import pathlib

import requests
from bs4 import BeautifulSoup
import zipfile
import xml.etree.ElementTree as ET
import docx
from langchain_community.document_loaders import PyPDFLoader, Docx2txtLoader

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


def _load_docx(path: pathlib.Path) -> str:
    """Extract text from a .docx file reliably.

    Handles non-standard document part names (e.g. 'word/document2.xml')
    produced by Word Online / co-authoring revisions, extracts tables,
    and preserves heading levels as Markdown hashes for the chunker.
    """
    if not zipfile.is_zipfile(str(path)):
        with open(path, "rb") as f:
            header = f.read(8)
        if header.startswith(b"{\\rtf"):
            raise ValueError(f"The file '{path.name}' appears to be an RTF document. Please save it as .docx or .pdf.")
        if header.startswith(b"\xd0\xcf\x11\xe0"):
            raise ValueError(f"The file '{path.name}' is an older Word 97-2003 (.doc) file. Please save it as .docx or .pdf.")
        raise ValueError(f"The file '{path.name}' is not a valid DOCX file.")

    # 1. Primary parser: python-docx (resolves OPC relationships and structures)
    try:
        doc = docx.Document(str(path))
        parts = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            style_name = p.style.name if p.style else ""
            if style_name.startswith("Heading"):
                try:
                    lvl = int(style_name.split()[-1])
                    h = "#" * min(max(lvl, 1), 6)
                    parts.append(f"\n{h} {text}\n")
                except (ValueError, IndexError):
                    parts.append(f"\n# {text}\n")
            elif style_name == "Title":
                parts.append(f"\n# {text}\n")
            elif style_name == "Subtitle":
                parts.append(f"\n## {text}\n")
            else:
                parts.append(text)

        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    table_rows.append(" | ".join(cells))
            if table_rows:
                parts.append("\n" + "\n".join(table_rows) + "\n")

        extracted = "\n\n".join(parts).strip()
        if extracted:
            return extracted
    except Exception:
        pass

    # 2. Fallback: Parse XML directly from zip archive (handles any word/document*.xml)
    try:
        texts = []
        with zipfile.ZipFile(str(path), "r") as z:
            doc_xmls = [n for n in z.namelist() if n.startswith("word/document") and n.endswith(".xml")]
            # Sort so word/document.xml or document2.xml are processed in order
            doc_xmls.sort()
            for name in doc_xmls:
                tree = ET.fromstring(z.read(name))
                for p in tree.iter():
                    if p.tag.endswith("}p"):
                        p_text = "".join(n.text for n in p.iter() if n.tag.endswith("}t") and n.text)
                        if p_text.strip():
                            texts.append(p_text.strip())
        extracted = "\n\n".join(texts).strip()
        if extracted:
            return extracted
    except Exception:
        pass

    # 3. Last fallback: Docx2txtLoader
    loader = Docx2txtLoader(str(path))
    docs = loader.load()
    return "\n\n".join(doc.page_content for doc in docs)


def load_document(file_path: str) -> str:
    path = pathlib.Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        loader = PyPDFLoader(str(path))
        docs = loader.load()
        return "\n\n".join(doc.page_content for doc in docs)
    elif suffix == ".docx":
        return _load_docx(path)
    elif suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: {path.suffix}")


def load_url(url: str) -> tuple[str, str]:
    """Fetch a web page and return (title, extracted plain text)."""
    response = requests.get(
        url,
        timeout=20,
        headers={"User-Agent": "Mozilla/5.0 (compatible; rag-demo-ingestor)"},
    )
    response.raise_for_status()

    soup = BeautifulSoup(response.text, "html.parser")
    # Drop non-content elements before extracting text
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript", "iframe"]):
        tag.decompose()

    title = (soup.title.string or "").strip() if soup.title and soup.title.string else ""

    # Rewrite <h1>-<h6> as Markdown headings so the heading-aware chunker can
    # prefix each chunk with its section, same as for uploaded documents.
    for level in range(1, 7):
        for heading in soup.find_all(f"h{level}"):
            heading_text = heading.get_text(strip=True)
            if heading_text:
                heading.string = f"\n{'#' * level} {heading_text}\n"
    lines = (line.strip() for line in soup.get_text("\n").splitlines())
    text = "\n".join(line for line in lines if line)
    return title, text
