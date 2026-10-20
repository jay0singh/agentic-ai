import pytest

import core.ingestor as ingestor


def test_load_txt(tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("Plain text policy notes.", encoding="utf-8")
    assert ingestor.load_document(str(f)) == "Plain text policy notes."


def test_load_md(tmp_path):
    f = tmp_path / "guide.md"
    f.write_text("# Heading\n\nSome markdown body.", encoding="utf-8")
    assert "Some markdown body." in ingestor.load_document(str(f))


def test_unsupported_extension_raises(tmp_path):
    f = tmp_path / "data.csv"
    f.write_text("a,b,c", encoding="utf-8")
    with pytest.raises(ValueError):
        ingestor.load_document(str(f))


def test_missing_file_raises():
    with pytest.raises(FileNotFoundError):
        ingestor.load_document("does-not-exist.txt")


def _fake_url_response(monkeypatch, html):
    class FakeResponse:
        text = html
        def raise_for_status(self):
            pass

    monkeypatch.setattr(ingestor.requests, "get", lambda *a, **k: FakeResponse())


def test_load_url_extracts_title_and_text(monkeypatch):
    _fake_url_response(monkeypatch, """
    <html><head><title> Sample Page </title><script>ignored()</script></head>
    <body>
      <nav>menu junk</nav>
      <h1>Welcome</h1>
      <p>This is the real content.</p>
      <footer>footer junk</footer>
    </body></html>
    """)

    title, text = ingestor.load_url("https://example.com/page")
    assert title == "Sample Page"
    assert "This is the real content." in text
    assert "menu junk" not in text
    assert "footer junk" not in text
    assert "ignored()" not in text


def test_load_url_converts_html_headings_to_markdown(monkeypatch):
    _fake_url_response(monkeypatch, """
    <html><body>
      <h1>Design principles</h1>
      <p>We follow five principles.</p>
      <h2>Secure by default</h2>
      <p>Auth is built in.</p>
    </body></html>
    """)

    _, text = ingestor.load_url("https://example.com/page")
    assert "# Design principles" in text
    assert "## Secure by default" in text
    # And the chunker turns those into section-prefixed chunks:
    from core.chunker import chunk_text
    chunks = chunk_text(text)
    assert any(c.startswith("Secure by default > ") for c in chunks)


def test_load_docx_standard(tmp_path):
    import docx
    doc_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph("This describes the system architecture.")
    doc.save(str(doc_path))

    text = ingestor.load_document(str(doc_path))
    assert "# Architecture Overview" in text
    assert "This describes the system architecture." in text


def test_load_docx_with_document2_xml(tmp_path):
    """Simulates documents where Word / co-authoring saves the main document as word/document2.xml."""
    import zipfile
    doc_path = tmp_path / "document2_variant.docx"
    with zipfile.ZipFile(str(doc_path), "w") as z:
        xml_content = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:body><w:p><w:r><w:t>Hello from document2.xml</w:t></w:r></w:p></w:body>'
            '</w:document>'
        )
        z.writestr("word/document2.xml", xml_content)

    text = ingestor.load_document(str(doc_path))
    assert "Hello from document2.xml" in text


def test_load_docx_invalid_file(tmp_path):
    f = tmp_path / "fake.docx"
    f.write_bytes(b"not a zip file content")
    with pytest.raises(ValueError, match="not a valid DOCX file"):
        ingestor.load_document(str(f))
